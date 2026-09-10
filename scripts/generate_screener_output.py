from pathlib import Path
import sys
import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter


# ================================================================
# PROJECT ROOT
# ================================================================

ROOT = Path(__file__).resolve().parents[1]

if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


# ================================================================
# IMPORT EXISTING SCREENER ENGINE
# ================================================================

from src.screener.engine import (
    load_config,
    load_ratio_data,
    load_supporting_data,
    build_screener_dataset,
    apply_preset,
)


# ================================================================
# PATHS
# ================================================================

CONFIG_PATH = ROOT / "config" / "screener_config.yaml"

OUTPUT_DIR = ROOT / "output"

OUTPUT_PATH = OUTPUT_DIR / "screener_output.xlsx"


# ================================================================
# KPI COLUMNS REQUIRED BY DAY 17
# ================================================================

KPI_COLUMNS = [
    "return_on_equity_pct",
    "roce_pct",
    "net_profit_margin_pct",
    "operating_profit_margin_pct",
    "debt_to_equity",
    "interest_coverage",
    "free_cash_flow_cr",
    "cfo_pat_ratio",
    "fcf_positive_score",
    "revenue_cagr_3yr",
    "revenue_cagr_5yr",
    "pat_cagr_5yr",
    "eps_cagr_5yr",
    "asset_turnover",
    "sales",
    "net_profit",
    "eps",
    "dividend_yield_pct",
    "dividend_payout_ratio_pct",
    "composite_quality_score",
]


# ================================================================
# EXCEL SHEET NAME CLEANER
# ================================================================

def clean_sheet_name(name):

    invalid_characters = [
        "\\",
        "/",
        "*",
        "?",
        ":",
        "[",
        "]",
    ]

    result = str(name)

    for character in invalid_characters:
        result = result.replace(character, "_")

    return result[:31]


# ================================================================
# EXCEL FORMATTING
# ================================================================

def format_workbook(workbook):

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    green_fill = PatternFill(
        fill_type="solid",
        fgColor="C6EFCE",
    )

    red_fill = PatternFill(
        fill_type="solid",
        fgColor="FFC7CE",
    )

    percentage_columns = {
        "return_on_equity_pct",
        "roce_pct",
        "net_profit_margin_pct",
        "operating_profit_margin_pct",
        "revenue_cagr_3yr",
        "revenue_cagr_5yr",
        "pat_cagr_5yr",
        "eps_cagr_5yr",
        "dividend_yield_pct",
        "dividend_payout_ratio_pct",
        "composite_quality_score",
    }

    decimal_columns = {
        "debt_to_equity",
        "interest_coverage",
        "asset_turnover",
        "cfo_pat_ratio",
    }

    money_columns = {
        "free_cash_flow_cr",
        "sales",
        "net_profit",
        "eps",
    }

    for worksheet in workbook.worksheets:

        # --------------------------------------------------------
        # Freeze first row
        # --------------------------------------------------------

        worksheet.freeze_panes = "A2"

        # --------------------------------------------------------
        # Header formatting
        # --------------------------------------------------------

        for cell in worksheet[1]:

            cell.fill = header_fill

            cell.font = header_font

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        # --------------------------------------------------------
        # Auto filter
        # --------------------------------------------------------

        worksheet.auto_filter.ref = worksheet.dimensions

        # --------------------------------------------------------
        # Find headers
        # --------------------------------------------------------

        headers = {
            cell.column: cell.value
            for cell in worksheet[1]
        }

        # --------------------------------------------------------
        # Number formats
        # --------------------------------------------------------

        for column_number, header in headers.items():

            if header in percentage_columns:

                for row_number in range(
                    2,
                    worksheet.max_row + 1,
                ):

                    worksheet.cell(
                        row=row_number,
                        column=column_number,
                    ).number_format = "0.00"

            elif header in decimal_columns:

                for row_number in range(
                    2,
                    worksheet.max_row + 1,
                ):

                    worksheet.cell(
                        row=row_number,
                        column=column_number,
                    ).number_format = "0.00"

            elif header in money_columns:

                for row_number in range(
                    2,
                    worksheet.max_row + 1,
                ):

                    worksheet.cell(
                        row=row_number,
                        column=column_number,
                    ).number_format = "#,##0.00"

        # --------------------------------------------------------
        # KPI cell highlighting
        # --------------------------------------------------------

        for column_number, header in headers.items():

            if header not in KPI_COLUMNS:
                continue

            for row_number in range(
                2,
                worksheet.max_row + 1,
            ):

                cell = worksheet.cell(
                    row=row_number,
                    column=column_number,
                )

                value = cell.value

                if value is None:
                    continue

                if not isinstance(value, (int, float)):
                    continue

                if pd.isna(value):
                    continue

                if value > 0:

                    cell.fill = green_fill

                elif value < 0:

                    cell.fill = red_fill

        # --------------------------------------------------------
        # Column widths
        # --------------------------------------------------------

        for column_cells in worksheet.columns:

            column_letter = get_column_letter(
                column_cells[0].column
            )

            maximum_length = 0

            for cell in column_cells:

                try:
                    cell_length = len(
                        str(cell.value)
                    )
                except Exception:
                    cell_length = 0

                maximum_length = max(
                    maximum_length,
                    cell_length,
                )

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                maximum_length + 2,
                30,
            )


# ================================================================
# MAIN
# ================================================================

def main():

    print("=" * 70)

    print(
        "NIFTY 100 - SPRINT 3 DAY 17"
    )

    print(
        "SCREENER OUTPUT GENERATOR"
    )

    print("=" * 70)

    # ------------------------------------------------------------
    # OUTPUT DIRECTORY
    # ------------------------------------------------------------

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    # ------------------------------------------------------------
    # STEP 1
    # ------------------------------------------------------------

    print(
        "\n[1] Loading configuration..."
    )

    config = load_config(
        CONFIG_PATH
    )

    print(
        "[OK] Configuration loaded"
    )

    # ------------------------------------------------------------
    # STEP 2
    # ------------------------------------------------------------

    print(
        "\n[2] Loading financial ratio data..."
    )

    ratios = load_ratio_data()

    print(
        f"[OK] financial_ratios rows: "
        f"{len(ratios):,}"
    )

    # ------------------------------------------------------------
    # STEP 3
    # ------------------------------------------------------------

    print(
        "\n[3] Loading supporting source data..."
    )

    sources = load_supporting_data()

    for name, dataframe in sources.items():

        print(
            f"[OK] {name:<15} "
            f"rows={len(dataframe):,}"
        )

    # ------------------------------------------------------------
    # STEP 4
    # ------------------------------------------------------------

    print(
        "\n[4] Building screener dataset..."
    )

    dataset = build_screener_dataset(
        ratios,
        sources,
    )

    print(
        f"[OK] Screener dataset rows: "
        f"{len(dataset)}"
    )

    print(
        f"[OK] Unique companies: "
        f"{dataset['company_id'].nunique()}"
    )

    # ------------------------------------------------------------
    # DATA QUALITY CHECKS
    # ------------------------------------------------------------

    if "company_name" in dataset.columns:

        missing_names = (
            dataset["company_name"]
            .isna()
            .sum()
        )

        print(
            f"[CHECK] Missing company names: "
            f"{missing_names}"
        )

    if "broad_sector" in dataset.columns:

        missing_sectors = (
            dataset["broad_sector"]
            .isna()
            .sum()
        )

        print(
            f"[CHECK] Missing sectors: "
            f"{missing_sectors}"
        )

    # ------------------------------------------------------------
    # STEP 5
    # ------------------------------------------------------------

    print(
        "\n[5] Validating composite quality score..."
    )

    if (
        "composite_quality_score"
        not in dataset.columns
    ):

        raise RuntimeError(
            "composite_quality_score "
            "is missing from dataset."
        )

    scores = pd.to_numeric(
        dataset[
            "composite_quality_score"
        ],
        errors="coerce",
    )

    valid_scores = scores.dropna()

    if valid_scores.empty:

        raise RuntimeError(
            "No valid composite quality "
            "scores were generated."
        )

    minimum_score = valid_scores.min()

    maximum_score = valid_scores.max()

    print(
        f"[CHECK] Minimum score: "
        f"{minimum_score:.2f}"
    )

    print(
        f"[CHECK] Maximum score: "
        f"{maximum_score:.2f}"
    )

    if minimum_score < 0:

        raise RuntimeError(
            "Composite score below 0."
        )

    if maximum_score > 100:

        raise RuntimeError(
            "Composite score above 100."
        )

    print(
        "[OK] Composite quality score "
        "is within 0-100"
    )

    # ------------------------------------------------------------
    # STEP 6
    # ------------------------------------------------------------

    print(
        "\n[6] Generating preset results..."
    )

    presets = config.get(
        "presets",
        {},
    )

    if not presets:

        raise RuntimeError(
            "No presets found in "
            "screener_config.yaml"
        )

    generated_results = {}

    # ------------------------------------------------------------
    # EACH PRESET
    # ------------------------------------------------------------

    for preset_name in presets:

        print(
            f"\n[PRESET] {preset_name}"
        )

        result = apply_preset(
            dataset.copy(),
            preset_name,
            config,
            ratios,
        )

        if result is None:

            result = dataset.iloc[
                0:0
            ].copy()

        # --------------------------------------------------------
        # Identity columns
        # --------------------------------------------------------

        identity_columns = [
            column
            for column in [
                "company_id",
                "company_name",
                "broad_sector",
                "sub_sector",
            ]
            if column in result.columns
        ]

        # --------------------------------------------------------
        # KPI columns actually present
        # --------------------------------------------------------

        available_kpis = [
            column
            for column in KPI_COLUMNS
            if column in result.columns
        ]

        # --------------------------------------------------------
        # Final output columns
        # --------------------------------------------------------

        output_columns = (
            identity_columns
            + [
                column
                for column in available_kpis
                if column not in identity_columns
            ]
        )

        output = result[
            output_columns
        ].copy()

        # --------------------------------------------------------
        # Sort by composite score
        # --------------------------------------------------------

        if (
            "composite_quality_score"
            in output.columns
        ):

            output = output.sort_values(
                by="composite_quality_score",
                ascending=False,
                na_position="last",
            )

        output = output.reset_index(
            drop=True
        )

        generated_results[
            preset_name
        ] = output

        print(
            f"[OK] {len(output)} companies"
        )

    # ------------------------------------------------------------
    # STEP 7
    # ------------------------------------------------------------

    print(
        "\n[7] Writing screener_output.xlsx..."
    )

    with pd.ExcelWriter(
        OUTPUT_PATH,
        engine="openpyxl",
    ) as writer:

        for (
            preset_name,
            dataframe,
        ) in generated_results.items():

            sheet_name = clean_sheet_name(
                preset_name
            )

            dataframe.to_excel(
                writer,
                sheet_name=sheet_name,
                index=False,
            )

    print(
        "[OK] Excel workbook written"
    )

    # ------------------------------------------------------------
    # STEP 8
    # ------------------------------------------------------------

    print(
        "\n[8] Applying Excel formatting..."
    )

    workbook = load_workbook(
        OUTPUT_PATH
    )

    format_workbook(
        workbook
    )

    workbook.save(
        OUTPUT_PATH
    )

    print(
        "[OK] Formatting complete"
    )

    # ------------------------------------------------------------
    # STEP 9
    # ------------------------------------------------------------

    print(
        "\n[9] Validating workbook..."
    )

    check_workbook = load_workbook(
        OUTPUT_PATH,
        read_only=True,
    )

    sheet_names = (
        check_workbook.sheetnames
    )

    print(
        f"[CHECK] Sheets created: "
        f"{len(sheet_names)}"
    )

    for sheet_name in sheet_names:

        worksheet = (
            check_workbook[sheet_name]
        )

        company_count = (
            worksheet.max_row - 1
        )

        print(
            f"        {sheet_name:<28} "
            f"{company_count:>3} companies"
        )

    # ------------------------------------------------------------
    # EXPECT EXACTLY 6 PRESETS
    # ------------------------------------------------------------

    if len(sheet_names) != 6:

        raise RuntimeError(
            f"Expected 6 sheets, "
            f"found {len(sheet_names)}."
        )

    # ------------------------------------------------------------
    # FINAL MESSAGE
    # ------------------------------------------------------------

    print(
        "\n[OK] Workbook created successfully:"
    )

    print(
        f"     {OUTPUT_PATH}"
    )

    print("\n" + "=" * 70)

    print(
        "DAY 17 SCREENER OUTPUT COMPLETE"
    )

    print("=" * 70)


if __name__ == "__main__":
    main()